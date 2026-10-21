"""HR Attrition Dashboard builder — loads data/employees.csv and builds
HR_Attrition_Dashboard.xlsx (Dashboard sheet with KPI cards + native Excel
charts fed by the Summary sheet, plus the full Data sheet).

Run:  python build_dashboard.py
"""
import pandas as pd
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter

BASE = Path(__file__).parent
NAVY, GOLD = "1F3864", "D9A441"
df = pd.read_csv(BASE / "data" / "employees.csv")
left = df["Attrition"] == "Yes"
rate = left.mean() * 100
by_dept = df.groupby("Department")["Attrition"].apply(lambda s: (s == "Yes").mean() * 100).sort_values(ascending=False)
ot_rate = df.groupby("OverTime")["Attrition"].apply(lambda s: (s == "Yes").mean() * 100)
sat_rate = df.groupby("JobSatisfaction")["Attrition"].apply(lambda s: (s == "Yes").mean() * 100)

wb = Workbook()
ws = wb.active; ws.title = "Dashboard"
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws["B2"] = "HR ATTRITION DASHBOARD"
ws["B2"].font = Font(name="Calibri", size=18, bold=True, color=NAVY)
ws["B3"] = "Excel Dashboard  •  Synthetic HR dataset  •  1,500 employees"
ws["B3"].font = Font(name="Calibri", size=10, italic=True, color="808080")
kpis = [("TOTAL EMPLOYEES", f"{len(df):,}"), ("LEFT THE COMPANY", f"{int(left.sum()):,}"),
        ("ATTRITION RATE", f"{rate:.1f}%"), ("AVG MONTHLY INCOME", f"${df.MonthlyIncome.mean():,.0f}"),
        ("OVERTIME ATTRITION", f"{ot_rate['Yes']:.1f}%")]
GOLD_FILL = PatternFill("solid", fgColor=GOLD)
col = 2
for label, val in kpis:
    c1 = ws.cell(row=5, column=col, value=label)
    c1.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF"); c1.fill = GOLD_FILL
    c1.alignment = Alignment(horizontal="center")
    c2 = ws.cell(row=6, column=col, value=val)
    c2.font = Font(name="Calibri", size=16, bold=True, color=NAVY)
    c2.fill = PatternFill("solid", fgColor="EAF0F7"); c2.alignment = Alignment(horizontal="center")
    ws.column_dimensions[get_column_letter(col)].width = 22
    col += 2
sm = wb.create_sheet("Summary")
sm["A1"], sm["B1"] = "Department", "AttritionRatePct"
for i, (d, v) in enumerate(by_dept.items(), start=2):
    sm.cell(row=i, column=1, value=d); sm.cell(row=i, column=2, value=round(float(v), 1))
sm["D1"], sm["E1"] = "OverTime", "AttritionRatePct"
for i, (d, v) in enumerate(ot_rate.items(), start=2):
    sm.cell(row=i, column=4, value=d); sm.cell(row=i, column=5, value=round(float(v), 1))
sm["G1"], sm["H1"] = "JobSatisfaction", "AttritionRatePct"
for i, (d, v) in enumerate(sat_rate.items(), start=2):
    sm.cell(row=i, column=7, value=str(d)); sm.cell(row=i, column=8, value=round(float(v), 1))
for cats, vals, title, anchor in [
        (Reference(sm, min_col=1, min_row=2, max_row=1 + len(by_dept)),
         Reference(sm, min_col=2, min_row=1, max_row=1 + len(by_dept)),
         "Attrition Rate by Department (%)", "B9"),
        (Reference(sm, min_col=7, min_row=2, max_row=5),
         Reference(sm, min_col=8, min_row=1, max_row=5),
         "Attrition Rate by Job Satisfaction (1=Low, 4=High)", "B24")]:
    ch = BarChart(); ch.type = "col"; ch.title = title; ch.legend = None
    ch.add_data(vals, titles_from_data=True); ch.set_categories(cats)
    ch.width, ch.height = 15, 7.5
    ws.add_chart(ch, anchor)
ws2 = wb.create_sheet("Data"); ws2.append(list(df.columns))
for r in df.itertuples(index=False):
    ws2.append(list(r))
wb.save(BASE / "HR_Attrition_Dashboard.xlsx")
print(f"Built HR_Attrition_Dashboard.xlsx — attrition rate {rate:.1f}%")
